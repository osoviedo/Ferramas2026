"""Genera Informe_de_Pruebas_Ferramas.docx y Plan_de_Pruebas_Ferramas.docx."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1] / "docs"
FECHA = "19/07/2026"
VERSION = "1.1"
AUTORES = "Equipo Ferramas"


def _set_cell_shading(cell, hex_color: str) -> None:
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(
        qn("w:shd"),
        {
            qn("w:fill"): hex_color,
            qn("w:val"): "clear",
        },
    )
    shading.append(shd)


def _style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def _title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)


def _subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def _h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
        _set_cell_shading(hdr[i], "1A3A5C")
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
            for run in cells[c_idx].paragraphs[0].runs:
                run.font.size = Pt(10)
    doc.add_paragraph()


def build_informe() -> Path:
    doc = Document()
    _style_doc(doc)

    _title(doc, "Informe de Pruebas — Ferramas")
    _subtitle(doc, f"E-commerce de ferretería | Fecha: {FECHA} | Versión: {VERSION}")
    _subtitle(doc, f"Autores: {AUTORES}")
    doc.add_paragraph()

    _h1(doc, "1. Resumen ejecutivo")
    _p(
        doc,
        "Ferramas cuenta con una suite automatizada de 30 pruebas (unitarias, integración, "
        "mock, aceptación, carga y estrés). En la corrida del 19/07/2026 se obtuvo "
        "30/30 passed. La cobertura de código de app/ es 54%, un punto por debajo del "
        "objetivo documental (≥ 55%). No hay CI ni pruebas E2E de navegador.",
    )

    _h1(doc, "2. Contexto del sistema")
    _table(
        doc,
        ["Aspecto", "Detalle"],
        [
            ["Propósito", "E-commerce de ferretería (catálogo, carrito, pagos, roles)"],
            ["Stack", "Python 3.9+ / Flask 3, SQLAlchemy, Flask-Login, Jinja2, Bootstrap 5"],
            ["Base de datos", "SQLite local; MySQL vía DATABASE_URL"],
            ["Integraciones", "Mercado Pago, mindicador.cl, SMTP"],
            ["Deploy", "Render (render.yaml)"],
            ["Entrada", "python run.py"],
        ],
    )

    _h1(doc, "3. Ambiente y ejecución")
    _table(
        doc,
        ["Ítem", "Valor"],
        [
            ["Fecha de ejecución", FECHA],
            ["Python", "3.12.5"],
            ["Framework de pruebas", "pytest 9.0.3"],
            ["Plugins", "pytest-cov, pytest-html, pytest-mock, responses"],
            ["BD de prueba", "SQLite temporal (fixture tmp_path)"],
            ["MP_ACCESS_TOKEN", "TEST-123456789-abcdef (modo simulado)"],
            ["Comando", "python -m pytest -v"],
            ["Cobertura", "python -m pytest --cov=app --cov-report=term"],
        ],
    )

    _h1(doc, "4. Resultados globales")
    _table(
        doc,
        ["Métrica", "Resultado"],
        [
            ["Tests totales", "30"],
            ["Passed", "30"],
            ["Failed / Errors", "0 / 0"],
            ["Tiempo aproximado", "~2 min 8 s"],
            ["Cobertura app/", "54% (1005 stmts, 462 sin cubrir)"],
            ["Warnings", "388 (SQLAlchemy Query.get() legacy)"],
            ["CI/CD", "No configurado"],
            ["E2E navegador", "No existe"],
        ],
    )

    _h1(doc, "5. Resultados por tipo de prueba")
    _table(
        doc,
        ["Tipo", "Archivo", "Casos", "Estado"],
        [
            ["Unitarios", "tests/unit/test_models_and_auth.py", "5", "PASSED"],
            ["Mock", "tests/unit/test_services_mock.py", "4", "PASSED"],
            ["Integración", "tests/integration/test_api_and_flows.py", "10", "PASSED"],
            ["Aceptación CA01–CA07", "tests/acceptance/test_criterios_aceptacion.py", "7", "PASSED"],
            ["Carga", "tests/load/test_carga_api.py", "2", "PASSED"],
            ["Estrés", "tests/stress/test_estres_api.py", "2", "PASSED"],
        ],
    )

    _h1(doc, "6. Criterios de aceptación (CA01–CA07)")
    _table(
        doc,
        ["ID", "Criterio", "Resultado"],
        [
            ["CA01", "Cliente ve catálogo con precio y stock (API)", "PASSED"],
            ["CA02", "Filtro por categoría en web", "PASSED"],
            ["CA03", "Agregar producto al carrito autenticado", "PASSED"],
            ["CA04", "Login rechaza credenciales inválidas", "PASSED"],
            ["CA05", "API categorías coincide con BD", "PASSED"],
            ["CA06", "Imagen de producto accesible en /static/", "PASSED"],
            ["CA07", "Webhook Mercado Pago acepta payload de test", "PASSED"],
        ],
    )

    _h1(doc, "7. Cobertura por módulo")
    _table(
        doc,
        ["Módulo", "Cobertura", "Observación"],
        [
            ["models/", "87–100%", "Bien cubierto"],
            ["divisa_service.py", "95%", "Mock de API externa"],
            ["auth_service.py", "83%", "Falta cambiar_password"],
            ["pago_service.py", "43%", "Solo modo simulado y preferencia básica"],
            ["email_service.py", "23%", "Sin tests"],
            ["views/tienda.py", "27%", "Checkout/pago casi sin cubrir"],
            ["views/carrito.py", "44%", "Solo agregar + checkout"],
            ["views/auth.py", "40%", "Solo login"],
            ["views/vendedor.py", "40%", "Sin tests"],
            ["views/bodeguero.py", "40%", "Sin tests"],
            ["views/admin.py", "36%", "Sin tests"],
            ["TOTAL app/", "54%", "Objetivo documental ≥ 55%"],
        ],
    )

    _h1(doc, "8. Alcance cubierto y no cubierto")
    _h2(doc, "8.1 Cubierto")
    _bullets(
        doc,
        [
            "API REST: /api/productos, /api/categorias, /api/divisas",
            "Login (éxito y fallo)",
            "Carrito: agregar producto y checkout → pedido pendiente",
            "Conversión USD con mock de mindicador.cl",
            "PagoService en modo simulado",
            "Rendimiento básico (50–200 peticiones secuenciales)",
        ],
    )
    _h2(doc, "8.2 No cubierto (automatizado)")
    _bullets(
        doc,
        [
            "Paneles vendedor / bodeguero / admin",
            "Flujo completo de pedido: aprobar → preparar → entregar",
            "Registro HTTP, perfil, logout",
            "Carrito: eliminar, actualizar, sincronizar",
            "Pago éxito/error y transferencia",
            "EmailService (SMTP)",
            "Suscripción newsletter",
            "Concurrencia real (load/stress son secuenciales)",
        ],
    )

    _h1(doc, "9. Riesgos y observaciones")
    _table(
        doc,
        ["Riesgo", "Severidad", "Mitigación / estado"],
        [
            ["API mindicador.cl caída", "Baja", "Mock + fallback 950/1050"],
            ["Token Mercado Pago inválido", "Media", "Modo simulado TEST"],
            ["Cobertura 54% < 55%", "Baja", "Ampliar tests en tienda/email/roles"],
            ["Sin CI en deploy", "Media", "Ejecución manual pre-deploy"],
            ["388 warnings Query.get()", "Baja", "Deprecación SQLAlchemy; no bloquea"],
            ["Corridas pytest en paralelo", "Media", "Race conditions en SQLite; ejecutar en serie"],
        ],
    )

    _h1(doc, "10. Registro de defectos")
    _p(
        doc,
        "Tras la corrida del 19/07/2026 no hay defectos funcionales abiertos en la suite "
        "automatizada (30/30 passed). Los errores observados en una corrida paralela de "
        "cobertura fueron de aislamiento de fixtures/SQLite, no del producto.",
    )
    _table(
        doc,
        ["ID", "Caso", "Módulo", "Descripción", "Severidad", "Estado"],
        [["—", "—", "—", "Sin defectos abiertos", "—", "Cerrado"]],
    )

    _h1(doc, "11. Conclusión")
    _p(
        doc,
        "La suite académica está sana y cumple los criterios de aceptación CA01–CA07. "
        "Para cerrar el criterio de cobertura (≥ 55%) y el flujo de negocio end-to-end, "
        "se recomienda ampliar pruebas de roles (vendedor/bodeguero/admin), registro y "
        "pago simulado, además de incorporar CI.",
    )

    out = DOCS / "Informe_de_Pruebas_Ferramas.docx"
    doc.save(out)
    return out


def build_plan() -> Path:
    doc = Document()
    _style_doc(doc)

    _title(doc, "Plan de Pruebas — Ferramas")
    _subtitle(doc, f"E-commerce de ferretería | Fecha: {FECHA} | Versión: {VERSION}")
    _subtitle(doc, f"Autores: {AUTORES}")
    doc.add_paragraph()

    _h1(doc, "1. Introducción")
    _h2(doc, "1.1 Resumen ejecutivo")
    _p(
        doc,
        "Ferramas es un e-commerce de ferretería desarrollado en Flask con API REST, "
        "SQLite y frontend Jinja2. Este plan define la estrategia de pruebas para validar "
        "API REST, flujo carrito–checkout–pedido, autenticación, pagos simulados y "
        "rendimiento. Criterio de éxito: 30+ tests passed y cobertura ≥ 55%.",
    )

    _h2(doc, "1.2 Objetivo")
    _p(
        doc,
        "Asegurar que Ferramas cumple los requisitos funcionales del e-commerce "
        "(catálogo, carrito, checkout, pagos, roles) antes de evaluación o despliegue.",
    )

    _h2(doc, "1.3 Alcance")
    _table(
        doc,
        ["En alcance", "Fuera de alcance (fase 1)"],
        [
            ["API REST pública", "UI visual / responsive detallado"],
            ["Auth (login, registro)", "Penetration testing"],
            ["Carrito y checkout", "Pruebas en producción con MP real"],
            ["Pago simulado + webhook mock", "E2E multi-navegador"],
            ["DivisaService y PagoService", "Email SMTP real en CI"],
            ["Rendimiento básico de API", ""],
        ],
    )

    _h2(doc, "1.4 Elementos a probar")
    _bullets(
        doc,
        [
            "API REST — app/api/producto_api.py (/api/productos, /api/categorias, /api/divisas, webhook MP)",
            "Vistas — app/views/tienda.py, carrito.py, auth.py, vendedor.py, bodeguero.py, admin.py",
            "Servicios — AuthService, DivisaService, PagoService, EmailService",
            "Modelos — Producto, Usuario, Carrito, Pedido",
            "Suite — carpeta tests/ (30 casos actuales)",
        ],
    )

    _h1(doc, "2. Riesgos")
    _table(
        doc,
        ["N°", "Riesgo", "Gravedad", "Acción"],
        [
            ["1", "API mindicador.cl caída", "Baja", "Mock + valores por defecto"],
            ["2", "Token Mercado Pago inválido", "Media", "Modo simulado TEST"],
            ["3", "Conflicto BD de prueba / paralelismo", "Media", "SQLite temporal; no correr pytest en paralelo"],
            ["4", "Cobertura bajo 55%", "Baja", "Ampliar tests Fase 2"],
            ["5", "Sin CI en deploy Render", "Media", "Workflow GitHub Actions propuesto"],
        ],
    )

    _h1(doc, "3. Estrategia de pruebas")
    _h2(doc, "3.1 Fase 1 — Automatizada (actual)")
    _table(
        doc,
        ["Tipo", "Herramienta", "Carpeta", "Comando"],
        [
            ["Unitarias", "pytest", "tests/unit/", "pytest -m unit -v"],
            ["Integración", "Flask test client", "tests/integration/", "pytest -m integration -v"],
            ["Mock", "unittest.mock + responses", "tests/unit/test_services_mock.py", "pytest -m mock -v"],
            ["Aceptación", "pytest CA01–CA07", "tests/acceptance/", "pytest -m acceptance -v"],
            ["Carga", "pytest (50 req)", "tests/load/", "pytest -m load -v"],
            ["Estrés", "pytest (200 req)", "tests/stress/", "pytest -m stress -v"],
            ["Todos", "pytest", "tests/", "python -m pytest -v"],
        ],
    )

    _h2(doc, "3.2 Fase 2 — Ampliación recomendada")
    _table(
        doc,
        ["Prioridad", "Caso", "Módulo", "Tipo"],
        [
            ["P1", "Vendedor: aprobar/rechazar pedido", "views/vendedor.py", "Integración"],
            ["P1", "Bodeguero: preparar/entregar", "views/bodeguero.py", "Integración"],
            ["P1", "Registro HTTP + carrito", "views/auth.py", "Integración"],
            ["P2", "Carrito: eliminar/actualizar", "views/carrito.py", "Integración"],
            ["P2", "Pago éxito/error (simulado)", "views/tienda.py", "Integración"],
            ["P2", "EmailService con SMTP mock", "email_service.py", "Unit + mock"],
            ["P3", "Admin: listar/crear usuarios", "views/admin.py", "Integración"],
            ["P3", "Webhook MP con estados", "pago_service.py", "Mock"],
        ],
    )

    _h2(doc, "3.3 Fase 3 — Manual / pre-demo")
    _table(
        doc,
        ["Caso", "Herramienta"],
        [
            ["Pago sandbox Mercado Pago", "scripts/demo_pago_mp.ps1 + ngrok"],
            ["API completa", "docs/postman_collection.json"],
            ["Demo exposición", "docs/GUIA_DEMO_EXPOSICION.md"],
        ],
    )

    _h2(doc, "3.4 Fase 4 — CI (propuesto)")
    _p(doc, "Workflow sugerido en .github/workflows/tests.yml:")
    _bullets(
        doc,
        [
            "pip install -r requirements.txt",
            "python -m pytest -v --cov=app --cov-fail-under=55",
        ],
    )

    _h1(doc, "4. Ambiente de pruebas")
    _table(
        doc,
        ["Variable", "Valor en tests"],
        [
            ["DATABASE_URL", "SQLite temporal (tmp_path)"],
            ["SECRET_KEY", "test-secret-key"],
            ["MP_ACCESS_TOKEN", "TEST-123456789-abcdef (modo simulado)"],
            ["Python", "3.9+ (probado en 3.12.5)"],
        ],
    )
    _p(doc, "Usuarios seed para pruebas manuales:")
    _table(
        doc,
        ["Rol", "Email", "Contraseña"],
        [
            ["Cliente", "cliente@ferramas.cl", "cliente123"],
            ["Vendedor", "vendedor@ferramas.cl", "vendedor123"],
            ["Bodeguero", "bodeguero@ferramas.cl", "bodeguero123"],
            ["Admin", "admin@ferramas.cl", "admin123"],
        ],
    )

    _h1(doc, "5. Casos de integración clave")
    _table(
        doc,
        ["ID", "Flujo", "Resultado esperado"],
        [
            ["CA-INT-01", "GET /api/productos", "200, JSON con id/nombre/precio/stock"],
            ["CA-INT-02", "Login → carrito → checkout", "Redirect /pago/iniciar/{id}, pedido pendiente"],
            ["CA-INT-03", "Mock mindicador → convertir_a=USD", "Precios en USD"],
            ["CA-INT-04*", "Checkout → vendedor aprueba", "Estado aprobado"],
            ["CA-INT-05*", "Aprobar → preparar → entregar", "Estado entregado"],
        ],
    )
    _p(doc, "* CA-INT-04 y CA-INT-05: propuestos en Fase 2 (aún no automatizados).")

    _h1(doc, "6. Criterios de aceptación (CA01–CA07)")
    _table(
        doc,
        ["ID", "Criterio", "Test automatizado"],
        [
            ["CA01", "Catálogo con precio y stock", "test_CA01_cliente_ve_catalogo_con_precio_y_stock"],
            ["CA02", "Filtro por categoría", "test_CA02_cliente_filtra_por_categoria"],
            ["CA03", "Agregar al carrito", "test_CA03_cliente_agrega_producto_al_carrito"],
            ["CA04", "Rechazo login inválido", "test_CA04_login_rechaza_credenciales_invalidas"],
            ["CA05", "API categorías = BD", "test_CA05_api_categorias_lista_todas_las_categorias"],
            ["CA06", "Imágenes /static/", "test_CA06_imagen_producto_accesible"],
            ["CA07", "Webhook MP test", "test_CA07_webhook_mercadopago_acepta_test"],
        ],
    )

    _h1(doc, "7. Criterios de entrada y salida")
    _h2(doc, "7.1 Entrada")
    _bullets(
        doc,
        [
            "pip install -r requirements.txt completado",
            "Python 3.9+ activo",
            "No ejecutar varias suites pytest concurrentes sobre el mismo entorno",
        ],
    )
    _h2(doc, "7.2 Salida (aprobado)")
    _bullets(
        doc,
        [
            "30+ tests passed, 0 failed",
            "Cobertura ≥ 55%",
            "CA01–CA07 verdes",
            "Load: 50 req en menos de 15 s",
            "Sin defectos abiertos de severidad alta",
        ],
    )
    _h2(doc, "7.3 Suspensión")
    _bullets(
        doc,
        [
            "Fallo de Mercado Pago real → usar modo simulado",
            "mindicador.cl caído → mock/fallback (ya implementado)",
        ],
    )

    _h1(doc, "8. Herramientas")
    _bullets(
        doc,
        [
            "pytest, pytest-cov, pytest-html, pytest-mock, responses",
            "Flask test client",
            "unittest.mock",
            "Postman (opcional, docs/postman_collection.json)",
            "python-docx (generación de este plan/informe)",
        ],
    )

    _h1(doc, "9. Cronograma sugerido")
    _table(
        doc,
        ["Actividad", "Duración", "Entregable"],
        [
            ["Ejecutar suite actual", "15 min", "Informe pytest / este informe Word"],
            ["Pruebas manuales Postman", "30 min", "Checklist API"],
            ["Demo pago MP sandbox", "20 min", "Flujo checkout completo"],
            ["Ampliar tests Fase 2 (P1)", "2–3 h", "+6–8 tests, cobertura > 60%"],
            ["Configurar CI", "1 h", ".github/workflows/tests.yml"],
        ],
    )

    _h1(doc, "10. Referencias")
    _bullets(
        doc,
        [
            "docs/CASOS_PRUEBA_FERRAMAS.md",
            "docs/Informe_de_Pruebas_Ferramas.docx (resultados de ejecución)",
            "README.md — sección Pruebas",
            "pytest.ini — markers unit/integration/mock/acceptance/load/stress",
        ],
    )

    out = DOCS / "Plan_de_Pruebas_Ferramas.docx"
    doc.save(out)
    return out


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    informe = build_informe()
    plan = build_plan()
    print(f"Generado: {informe}")
    print(f"Generado: {plan}")


if __name__ == "__main__":
    main()
