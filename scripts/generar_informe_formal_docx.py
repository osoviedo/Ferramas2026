"""Genera docs/Informe_Formal_Ferramas.docx para la entrega."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
BPMN_ASIS = DOCS / "bpmn" / "Ferremas_AS-IS.png"
BPMN_TOBE = DOCS / "bpmn" / "Ferremas_TO-BE.png"

FECHA = "19/07/2026"
VERSION = "1.0"
AUTORES = "Miguel Rojas, Jose Larenas, Oscar Oviedo"
ASIGNATURA = "Integración de Plataformas"
INSTITUCION = "Duoc UC"
REPO = "https://github.com/osoviedo/Ferramas2026"


def _set_cell_shading(cell, hex_color: str) -> None:
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(
        qn("w:shd"),
        {qn("w:fill"): hex_color, qn("w:val"): "clear"},
    )
    shading.append(shd)


def _style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def _center(doc: Document, text: str, size: int = 12, bold: bool = False, space_after: int = 6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def _h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def _h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _add_image(doc: Document, path: Path, width_inches: float = 6.2) -> None:
    if not path.exists():
        _p(doc, f"[Imagen no encontrada: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_cell_shading(hdr[i], "1A3A5C")
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
            for run in cells[c_idx].paragraphs[0].runs:
                run.font.size = Pt(10)
    doc.add_paragraph()


def _page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_portada(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    _center(doc, INSTITUCION, size=16, bold=True, space_after=4)
    _center(doc, ASIGNATURA, size=14, bold=False, space_after=24)
    _center(doc, "INFORME FORMAL DEL PROYECTO", size=18, bold=True, space_after=8)
    _center(doc, "Ferramas — E-commerce de Ferretería", size=16, bold=True, space_after=24)
    _center(doc, "Integración de plataformas web, API REST,", size=12)
    _center(doc, "pasarela de pagos y flujo de pedidos por roles", size=12, space_after=36)
    _center(doc, "Autores:", size=12, bold=True, space_after=4)
    for autor in ["Miguel Rojas", "Jose Larenas", "Oscar Oviedo"]:
        _center(doc, autor, size=12, space_after=2)
    doc.add_paragraph()
    _center(doc, f"Fecha: {FECHA}", size=12)
    _center(doc, f"Versión: {VERSION}", size=12)
    _center(doc, f"Repositorio: {REPO}", size=10, space_after=12)
    _page_break(doc)


def build_indice(doc: Document) -> None:
    _h1(doc, "Índice")
    items = [
        "1. Objetivos",
        "    1.1 Objetivo general",
        "    1.2 Objetivos específicos",
        "2. Problemática y soluciones propuestas",
        "    2.1 Problemática (AS-IS)",
        "    2.2 Solución propuesta (TO-BE)",
        "    2.3 Arquitectura e implementación",
        "3. Pruebas realizadas",
        "    3.1 Estrategia de pruebas",
        "    3.2 Resultados de la ejecución",
        "    3.3 Criterios de aceptación",
        "    3.4 Cobertura y gaps",
        "4. Conclusiones",
        "5. Anexos",
    ]
    for item in items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)
    _page_break(doc)


def build_objetivos(doc: Document) -> None:
    _h1(doc, "1. Objetivos")
    _h2(doc, "1.1 Objetivo general")
    _p(
        doc,
        "Desarrollar e integrar una plataforma e-commerce para una ferretería (Ferramas) "
        "que digitalice el ciclo de venta —desde el catálogo y el carrito hasta el pago, "
        "la preparación y la entrega— mediante una aplicación web Flask, API REST, "
        "pasarela de pagos (Mercado Pago) y paneles diferenciados por rol.",
    )
    _h2(doc, "1.2 Objetivos específicos")
    _bullets(
        doc,
        [
            "Modelar el proceso de negocio actual (AS-IS) y el proceso objetivo (TO-BE) con diagramas BPMN.",
            "Implementar catálogo de productos, carrito de compras y checkout con persistencia en base de datos.",
            "Integrar autenticación y autorización por roles: cliente, vendedor, bodeguero y administrador.",
            "Exponer una API REST pública para productos, categorías y divisas, con conversión CLP/USD/EUR.",
            "Integrar Mercado Pago (modo simulado y sandbox) para la gestión de pagos y webhooks.",
            "Automatizar pruebas unitarias, de integración, mock, aceptación, carga y estrés con pytest.",
            "Documentar el proyecto (README), el plan de pruebas y el informe de resultados para la entrega.",
        ],
    )


def build_problematica(doc: Document) -> None:
    _h1(doc, "2. Problemática y soluciones propuestas")

    _h2(doc, "2.1 Problemática (AS-IS)")
    _p(
        doc,
        "En el proceso AS-IS, la ferretería opera de forma predominantemente presencial y "
        "manual. El cliente solicita productos, el vendedor asesora, el bodeguero verifica "
        "stock y prepara pedidos, y el pago (efectivo, transferencia o tarjeta) se confirma "
        "con intervención del contador. Los reportes financieros y de ventas quedan a cargo "
        "del contador y del administrador. Este modelo presenta fricciones típicas:",
    )
    _bullets(
        doc,
        [
            "Dependencia de interacción presencial y tiempos de espera por verificación de stock.",
            "Riesgo de inconsistencias entre venta, stock y facturación al no haber un sistema unificado.",
            "Confirmación de pago y registro contable en pasos separados, con posible retrabajo si el pago falla.",
            "Visibilidad limitada del estado del pedido para el cliente.",
            "Reportes mensuales elaborados de forma reactiva, no como resultado automático del flujo digital.",
        ],
    )
    _p(doc, "Diagrama BPMN AS-IS (primera entrega):")
    _add_image(doc, BPMN_ASIS, width_inches=6.3)
    _caption(doc, "Figura 1. Diagrama BPMN AS-IS — Proceso actual de la ferretería (Bizagi Modeler).")

    _h2(doc, "2.2 Solución propuesta (TO-BE)")
    _p(
        doc,
        "El proceso TO-BE digitaliza e integra los mismos actores (cliente, vendedor, "
        "bodeguero, contador, administrador) sobre una plataforma web. El cliente consulta "
        "catálogo, arma carrito y paga en línea; el vendedor aprueba o rechaza pedidos; "
        "el bodeguero prepara y entrega; el sistema registra transacciones y facilita "
        "reportes. La solución reduce fricción, centraliza datos y permite trazabilidad "
        "del pedido (pendiente → aprobado → preparado → entregado).",
    )
    _bullets(
        doc,
        [
            "Canal digital 24/7 para catálogo, carrito y checkout.",
            "Integración de pagos con Mercado Pago (simulado en desarrollo; sandbox/producción configurable).",
            "Roles con paneles específicos que reflejan el BPMN TO-BE.",
            "API REST para consumo de productos, categorías y tasas de cambio (mindicador.cl).",
            "Suite de pruebas automatizadas que respaldan la calidad de la integración.",
        ],
    )
    _p(doc, "Diagrama BPMN TO-BE (primera entrega):")
    _add_image(doc, BPMN_TOBE, width_inches=6.3)
    _caption(doc, "Figura 2. Diagrama BPMN TO-BE — Proceso objetivo digitalizado (Bizagi Modeler).")

    _h2(doc, "2.3 Arquitectura e implementación")
    _p(
        doc,
        "Ferramas se implementó como aplicación monolítica modular en Flask, con vistas "
        "Jinja2/Bootstrap, modelos SQLAlchemy y servicios de dominio (auth, pago, divisa, email). "
        "La base de datos puede ser SQLite (desarrollo/pruebas) o MySQL (despliegue).",
    )
    _table(
        doc,
        ["Componente", "Tecnología / ubicación"],
        [
            ["Backend / web", "Flask 3, Flask-Login, Jinja2"],
            ["Persistencia", "SQLAlchemy + SQLite / MySQL"],
            ["API REST", "app/api/producto_api.py"],
            ["Pagos", "Mercado Pago SDK — app/services/pago_service.py"],
            ["Divisas", "mindicador.cl — app/services/divisa_service.py"],
            ["Frontend", "Bootstrap 5 + JS vanilla"],
            ["Despliegue", "Render (render.yaml)"],
            ["Repositorio", REPO],
        ],
    )
    _p(doc, "Flujo de pedido implementado (alineado al TO-BE):")
    _bullets(
        doc,
        [
            "Cliente: catálogo → carrito → checkout → pago (MP o transferencia).",
            "Vendedor: aprueba o rechaza pedidos pendientes.",
            "Bodeguero: prepara y marca entrega.",
            "Admin: gestión de usuarios y visión operativa.",
        ],
    )


def build_pruebas(doc: Document) -> None:
    _h1(doc, "3. Pruebas realizadas")

    _h2(doc, "3.1 Estrategia de pruebas")
    _p(
        doc,
        "Se definió un plan de pruebas con pytest que cubre unitarias, integración, mock "
        "de servicios externos, criterios de aceptación (CA01–CA07), carga y estrés. "
        "El detalle completo está en docs/Plan_de_Pruebas_Ferramas.docx.",
    )
    _table(
        doc,
        ["Tipo", "Carpeta", "Casos"],
        [
            ["Unitarias", "tests/unit/", "5"],
            ["Mock", "tests/unit/test_services_mock.py", "4"],
            ["Integración", "tests/integration/", "10"],
            ["Aceptación", "tests/acceptance/", "7"],
            ["Carga", "tests/load/", "2"],
            ["Estrés", "tests/stress/", "2"],
            ["Total", "tests/", "30"],
        ],
    )

    _h2(doc, "3.2 Resultados de la ejecución")
    _p(
        doc,
        f"Corrida del {FECHA} (Python 3.12.5, pytest 9.0.3), comando: python -m pytest -v.",
    )
    _table(
        doc,
        ["Métrica", "Resultado"],
        [
            ["Tests totales", "30"],
            ["Passed", "30"],
            ["Failed / Errors", "0 / 0"],
            ["Tiempo", "~2 min 8 s"],
            ["Cobertura app/", "54%"],
            ["Objetivo cobertura documental", "≥ 55%"],
            ["CI/CD", "No configurado"],
        ],
    )
    _p(
        doc,
        "El informe detallado de resultados se encuentra en docs/Informe_de_Pruebas_Ferramas.docx.",
    )

    _h2(doc, "3.3 Criterios de aceptación")
    _table(
        doc,
        ["ID", "Criterio", "Resultado"],
        [
            ["CA01", "Catálogo con precio y stock vía API", "PASSED"],
            ["CA02", "Filtro por categoría en web", "PASSED"],
            ["CA03", "Agregar al carrito autenticado", "PASSED"],
            ["CA04", "Rechazo de login inválido", "PASSED"],
            ["CA05", "API categorías = BD", "PASSED"],
            ["CA06", "Imágenes accesibles en /static/", "PASSED"],
            ["CA07", "Webhook Mercado Pago acepta test", "PASSED"],
        ],
    )

    _h2(doc, "3.4 Cobertura y gaps")
    _p(
        doc,
        "La cobertura global de app/ alcanzó 54%, levemente bajo el umbral documental del 55%. "
        "Los módulos con menor cobertura son email_service (23%), views/tienda (27%), "
        "paneles de roles (36–40%) y pago_service (43%). Se recomienda, como trabajo futuro, "
        "automatizar el flujo vendedor → bodeguero, registro HTTP y callbacks de pago.",
    )


def build_conclusiones(doc: Document) -> None:
    _h1(doc, "4. Conclusiones")
    _p(
        doc,
        "Ferramas cumple el objetivo de integrar una plataforma e-commerce alineada al "
        "proceso TO-BE modelado en BPMN: el cliente opera en canal digital, los roles "
        "operativos gestionan el ciclo del pedido y los pagos se integran vía Mercado Pago "
        "(con modo simulado seguro para desarrollo).",
    )
    _bullets(
        doc,
        [
            "Se entregó un producto funcional con README, repositorio y despliegue configurable en Render.",
            "Los diagramas AS-IS y TO-BE de la primera entrega explican la problemática y la solución.",
            "La suite de 30 pruebas automatizadas pasó completa (30/30) y valida los criterios CA01–CA07.",
            "Quedan como mejoras: subir cobertura sobre 55%, cubrir paneles por rol y pago end-to-end, e incorporar CI.",
        ],
    )
    _p(
        doc,
        "En síntesis, el proyecto demuestra integración de plataformas (web, API, pagos, "
        "divisas y roles) con evidencia de pruebas y documentación formal para evaluación.",
    )


def build_anexos(doc: Document) -> None:
    _h1(doc, "5. Anexos")
    _h2(doc, "Anexo A — Artefactos de entrega")
    _table(
        doc,
        ["Artefacto", "Ubicación"],
        [
            ["Diagrama BPMN AS-IS", "docs/bpmn/Ferremas_AS-IS.png"],
            ["Diagrama BPMN TO-BE", "docs/bpmn/Ferremas_TO-BE.png"],
            ["Plan de pruebas", "docs/Plan_de_Pruebas_Ferramas.docx"],
            ["Informe de pruebas", "docs/Informe_de_Pruebas_Ferramas.docx"],
            ["README del proyecto", "README.md"],
            ["Repositorio", REPO],
            ["Casos de prueba (guía)", "docs/CASOS_PRUEBA_FERRAMAS.md"],
            ["Colección Postman", "docs/postman_collection.json"],
        ],
    )

    _h2(doc, "Anexo B — Usuarios de prueba")
    _table(
        doc,
        ["Rol", "Email", "Contraseña"],
        [
            ["Administrador", "admin@ferramas.cl", "admin123"],
            ["Vendedor", "vendedor@ferramas.cl", "vendedor123"],
            ["Bodeguero", "bodeguero@ferramas.cl", "bodeguero123"],
            ["Cliente", "cliente@ferramas.cl", "cliente123"],
        ],
    )

    _h2(doc, "Anexo C — Cómo ejecutar el proyecto y las pruebas")
    _p(doc, "Proyecto:")
    _bullets(
        doc,
        [
            "pip install -r requirements.txt",
            "cp .env.example .env  (Windows: copy .env.example .env)",
            "python run.py",
        ],
    )
    _p(doc, "Pruebas:")
    _bullets(
        doc,
        [
            "python -m pytest -v",
            "python -m pytest --cov=app --cov-report=html",
        ],
    )

    _h2(doc, "Anexo D — Checklist de entrega")
    _table(
        doc,
        ["Requisito", "Estado"],
        [
            ["Diagramas BPMN (primera entrega)", "Incluidos (AS-IS y TO-BE)"],
            ["Proyecto desarrollado + repositorio + README", "Cumple"],
            ["Informe de pruebas + plan de pruebas", "Cumple (docs/)"],
            ["Informe formal (este documento)", "Cumple"],
        ],
    )


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _style_doc(doc)
    build_portada(doc)
    build_indice(doc)
    build_objetivos(doc)
    build_problematica(doc)
    build_pruebas(doc)
    build_conclusiones(doc)
    build_anexos(doc)

    out = DOCS / "Informe_Formal_Ferramas.docx"
    doc.save(out)
    print(f"Generado: {out}")


if __name__ == "__main__":
    main()
