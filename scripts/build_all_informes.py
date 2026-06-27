"""Genera los 4 informes Ferramas desde plantillas del escritorio."""
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

DESKTOP = Path(r"c:\Users\Jose\Desktop\test")
DOCS = Path(r"c:\Users\Jose\Desktop\Ferramas2026-main\docs")

PROJECT = "Ferramas"
AUTHORS = "Jose Larenas / Oscar Oviedo"
AUTHOR_1 = "Jose Larenas"
AUTHOR_2 = "Oscar Oviedo"
TEAM = "Equipo Ferramas (Jose Larenas, Oscar Oviedo)"
DATE = "26/06/2026"
DATE_ISO = "2026-06-26"
ORG = "Duoc UC"

# --- Rutas plantillas ---
TPL_INT = DESKTOP / "3.1.4 Plantilla Casos de prueba Integracion.xlsx"
TPL_DEF = DESKTOP / "3.3.4 Planilla_Registro_de_Defectos_ejemplo.xlsx"
TPL_IMP = DESKTOP / "3.4.4 Plantilla_Plan_de_Implantacion.docx"

OUT_INT = DOCS / "Casos_Integracion_Ferramas.xlsx"
OUT_DEF = DOCS / "Registro_Defectos_Ferramas.xlsx"
OUT_IMP = DOCS / "Plan_Implantacion_Ferramas.docx"


def _black_cell(ws, row, col, value):
    ws.cell(row, col, value)


def build_integracion():
    shutil.copy2(TPL_INT, OUT_INT)
    wb = openpyxl.load_workbook(OUT_INT)
    hc = wb["Hoja_de_Control"]
    ws = wb[[n for n in wb.sheetnames if "Integraci" in n][0]]

    # Hoja de control
    hc["B2"] = PROJECT
    hc["B3"] = ORG
    hc["C10"] = ORG
    hc["C11"] = PROJECT
    hc["C13"] = AUTHORS
    hc["F14"] = DATE
    hc["F21"] = DATE
    hc["D21"] = AUTHORS

    # Encabezado hoja pruebas
    ws["C2"] = PROJECT

    # Resumen casos (filas 13-14)
    cases_summary = [
        (
            "CA-INT-01",
            "API REST - SQLAlchemy",
            "GET /api/productos devuelve JSON desde SQLite",
            "pytest instalado; app Flask configurada",
        ),
        (
            "CA-INT-02",
            "views/carrito - models/pedido",
            "Usuario logueado agrega producto y crea pedido pendiente",
            "Usuario cliente@ferramas.cl / cliente123",
        ),
    ]
    for idx, (cid, comp, desc, pre) in enumerate(cases_summary):
        r = 13 + idx
        ws.cell(r, 2, cid)
        ws.cell(r, 4, comp)
        ws.cell(r, 5, desc)
        ws.cell(r, 6, pre)

    # Bloque detallado CA-INT-01 (fila 16+)
    ws.cell(16, 2, "CA-INT-01")
    ws.cell(16, 4, "API REST - SQLAlchemy")
    ws.cell(16, 5, "Validar listado de productos vía API")
    ws.cell(16, 6, "Servidor Flask o pytest client")

    steps_01 = [
        (1, "Ejecutar GET /api/productos", "URL /api/productos", "HTTP 200, JSON lista", "Sí"),
        (2, "Verificar campos del primer producto", "—", "id, nombre, precio, stock, categoria", "Sí"),
        (3, "Ejecutar test automatizado", "pytest test_get_api_productos_retorna_lista_json", "Test passed", "Sí"),
    ]
    for i, step in enumerate(steps_01):
        r = 18 + i
        ws.cell(r, 2, step[0])
        ws.cell(r, 3, step[1])
        ws.cell(r, 4, step[2])
        ws.cell(r, 5, step[3])
        ws.cell(r, 6, step[4])

    # Bloque CA-INT-02 (fila 25+)
    ws.cell(25, 2, "CA-INT-02")
    ws.cell(25, 4, "carrito - pedido")
    ws.cell(25, 5, "Flujo agregar al carrito y checkout")
    ws.cell(25, 6, "Sesión cliente autenticada")

    steps_02 = [
        (1, "Login POST /login", "cliente@ferramas.cl / cliente123", "Sesión activa", "Sí"),
        (2, "POST /carrito/agregar", '{"producto_id":1,"cantidad":2}', '{"ok":true}', "Sí"),
        (3, "POST /carrito/checkout", "modo_entrega=retiro", "Redirect /pago/iniciar/", "Sí"),
        (4, "Verificar pedido en BD", "último pedido", "estado=pendiente, total>0", "Sí"),
    ]
    for i, step in enumerate(steps_02):
        r = 27 + i
        ws.cell(r, 2, step[0])
        ws.cell(r, 3, step[1])
        ws.cell(r, 4, step[2])
        ws.cell(r, 5, step[3])
        ws.cell(r, 6, step[4])

    ws.cell(34, 2, "Se incluyen fichas CA-INT-01 y CA-INT-02 (CA-INT-03 divisas/mock documentado en CASOS_PRUEBA_FERRAMAS.md).")

    # CA-INT-03 insert after row 34 - append rows
    ws.insert_rows(35, 10)
    ws.cell(35, 2, "CA-INT-03")
    ws.cell(35, 4, "DivisaService - mindicador.cl (mock)")
    ws.cell(35, 5, "Conversión USD con API externa simulada")
    ws.cell(35, 6, "unittest.mock requests.get")
    ws.cell(36, 2, "Paso")
    ws.cell(36, 3, "Descripción de pasos a seguir")
    ws.cell(36, 4, "Datos Entrada")
    ws.cell(36, 5, "Salida Esperada")
    ws.cell(36, 6, "¿OK?")
    steps_03 = [
        (1, "Mock requests.get con dolar=1000", "mock divisas", "DivisaService retorna tasa", "Sí"),
        (2, "GET /api/productos?convertir_a=USD", "—", "moneda=USD en JSON", "Sí"),
    ]
    for i, step in enumerate(steps_03):
        r = 37 + i
        ws.cell(r, 2, step[0])
        ws.cell(r, 3, step[1])
        ws.cell(r, 4, step[2])
        ws.cell(r, 5, step[3])
        ws.cell(r, 6, step[4])

    wb.save(OUT_INT)
    shutil.copy2(OUT_INT, DESKTOP / "Casos_Integracion_FERRAMAS.xlsx")
    print("OK", OUT_INT)


def build_defectos():
    shutil.copy2(TPL_DEF, OUT_DEF)
    wb = openpyxl.load_workbook(OUT_DEF)
    ws = wb["Registro Defectos"]

    ws["E4"] = PROJECT
    ws["E5"] = AUTHOR_1
    ws["E6"] = AUTHOR_2
    ws["E7"] = TEAM

    # Limpiar ejemplos hotel
    for r in (10, 11):
        for c in range(1, 11):
            ws.cell(r, c, None)

    # Registro sin defectos abiertos tras pruebas
    ws.cell(10, 1, "—")
    ws.cell(10, 2, "Suite pytest")
    ws.cell(10, 3, 1)
    ws.cell(10, 4, datetime(2026, 6, 26))
    ws.cell(10, 5, "General")
    ws.cell(10, 6, "Sin defectos abiertos tras ejecutar 30 tests automatizados (26/06/2026).")
    ws.cell(10, 7, "Funcionalidad")
    ws.cell(10, 8, "Observación")
    ws.cell(10, 9, "Cerrado")
    ws.cell(10, 10, "Corregido fixture BD en tests; suite completa OK.")

    wb.save(OUT_DEF)
    shutil.copy2(OUT_DEF, DESKTOP / "Registro_Defectos_FERRAMAS.xlsx")
    print("OK", OUT_DEF)


def run_color(run):
    if run.font.color and run.font.color.rgb:
        c = run.font.color.rgb
        return (c[0], c[1], c[2])
    return None


def is_blue_run(run):
    c = run_color(run)
    if c in ((0, 0, 255), (0, 176, 80)):
        return True
    rpr = run._element.find(qn("w:rPr"))
    if rpr is not None:
        style = rpr.find(qn("w:rStyle"))
        if style is not None and "Info" in (style.get(qn("w:val")) or ""):
            return True
    return False


def set_black(run, text=None):
    if text is not None:
        run.text = text
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.italic = False


def paragraph_text(p):
    return "".join(r.text for r in p.runs).strip()


def clean_blue_runs(doc):
    for p in doc.paragraphs:
        for run in p.runs:
            if is_blue_run(run):
                run.text = ""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if is_blue_run(run):
                            run.text = ""


def replace_in_paragraphs(doc, mapping):
    for p in doc.paragraphs:
        t = paragraph_text(p)
        for old, new in mapping.items():
            if old in t or t == old:
                p.clear()
                r = p.add_run(new if t == old or old in t else t.replace(old, new))
                r.font.color.rgb = RGBColor(0, 0, 0)


def build_implantacion():
    shutil.copy2(TPL_IMP, OUT_IMP)
    doc = Document(OUT_IMP)

    mapping = {
        "Proyecto xxx": f"Proyecto {PROJECT}",
        "xxx": PROJECT,
        "<nombre de cliente>": ORG,
    }
    for p in doc.paragraphs:
        for run in p.runs:
            txt = run.text
            for old, new in mapping.items():
                if old in txt:
                    set_black(run, txt.replace(old, new))

    clean_blue_runs(doc)

    replacements = {
        "El objetivo general de la implementación del sistema": (
            f"El objetivo general es implantar {PROJECT}, e-commerce de ferretería, "
            "en ambiente local de desarrollo y demostración académica, permitiendo "
            "a usuarios consultar catálogo, realizar compras y pagar vía Mercado Pago sandbox."
        ),
        "Definir las actividades necesarias para la implantación del software.": (
            "Definir actividades para desplegar Ferramas en localhost, configurar variables "
            "de entorno (.env), base SQLite, ejecutar suite de pruebas y capacitar al equipo."
        ),
    }

    for p in doc.paragraphs:
        t = paragraph_text(p)
        for key, val in replacements.items():
            if key in t:
                p.clear()
                r = p.add_run(val)
                r.font.color.rgb = RGBColor(0, 0, 0)
                break

    # Actividades numeradas - buscar párrafos cortos de actividad y rellenar si vacíos
    activity_text = {
        "Contratación del Personal del Proyecto": (
            "Equipo académico Duoc UC: Jose Larenas (desarrollo) y Oscar Oviedo (integración MP y pruebas)."
        ),
        "Difusión de Propuesta en la Empresa": (
            "Presentación del prototipo Ferramas al docente y compañeros en evaluación presencial."
        ),
        "Adquisición de Equipo y materiales": (
            "PC con Windows 10+, Python 3.14, Git, editor VS Code/Cursor. Sin hardware adicional."
        ),
        "Capacitación": (
            "Capacitación interna: uso de python run.py, pytest, endpoints /api/productos y flujo de compra."
        ),
        "Entrega Preliminar De Modelo": (
            "Entrega repositorio GitHub https://github.com/Fitzroyal/Ferramas con documentación en docs/."
        ),
        "Prueba Preliminar": (
            "Ejecución de 30 tests pytest + demo manual en http://127.0.0.1:5000 el 26/06/2026."
        ),
        "Retroalimentación de las Pruebas Preliminares:": (
            "Retroalimentación del docente incorporada: consolidación en carpeta app/, suite de tests completa."
        ),
        "Liberación de Recursos y Entrega del Proyecto:": (
            "Cierre de sprint académico: tag en GitHub, informes 3.1.4–3.4.4 y marcha blanca local."
        ),
    }

    for i, p in enumerate(doc.paragraphs):
        t = paragraph_text(p)
        for title, content in activity_text.items():
            if t.strip() == title.strip():
                if i + 1 < len(doc.paragraphs):
                    nxt = doc.paragraphs[i + 1]
                    if len(paragraph_text(nxt)) < 20 or "Esta etapa" in paragraph_text(nxt):
                        nxt.clear()
                        r = nxt.add_run(content)
                        r.font.color.rgb = RGBColor(0, 0, 0)

    # Tabla cronograma si existe - primera tabla con Actividad
    for table in doc.tables:
        header = "".join(c.text for c in table.rows[0].cells)
        if "Actividad" in header or "Tiempo" in header:
            rows_data = [
                ("A", "Configuración ambiente Python/Flask", "1", "—"),
                ("B", "Ejecución suite de pruebas", "1", "A"),
                ("C", "Demo evaluación y documentación", "1", "B"),
            ]
            for ri, rowdata in enumerate(rows_data, start=1):
                if ri < len(table.rows):
                    for ci, val in enumerate(rowdata):
                        if ci < len(table.rows[ri].cells):
                            cell = table.rows[ri].cells[ci]
                            cell.paragraphs[0].clear()
                            r = cell.paragraphs[0].add_run(str(val))
                            r.font.color.rgb = RGBColor(0, 0, 0)
            break

    doc.save(OUT_IMP)
    shutil.copy2(OUT_IMP, DESKTOP / "Plan_Implantacion_FERRAMAS.docx")
    print("OK", OUT_IMP)


def main():
    build_integracion()
    build_defectos()
    build_implantacion()
    # Plan pruebas 3.2.4
    import subprocess
    subprocess.run(
        ["python", str(Path(__file__).parent / "build_plan_pruebas_docx.py")],
        check=True,
    )
    print("\nInformes generados en docs/ y Desktop/test/")
    print("  1. Casos_Integracion_FERRAMAS.xlsx  (3.1.4)")
    print("  2. Plan_de_Pruebas_FERRAMAS.docx    (3.2.4)")
    print("  3. Registro_Defectos_FERRAMAS.xlsx  (3.3.4)")
    print("  4. Plan_Implantacion_FERRAMAS.docx  (3.4.4)")


if __name__ == "__main__":
    main()
